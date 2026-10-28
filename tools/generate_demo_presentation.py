from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Student_Task_Manager_Project_Presentation.pptx"
UI_IMAGE = ROOT / "evidence" / "student-task-manager-ui.png"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

NAVY = RGBColor(18, 35, 31)
DEEP_NAVY = RGBColor(12, 27, 24)
GREEN = RGBColor(39, 126, 99)
MINT = RGBColor(188, 236, 221)
PALE_MINT = RGBColor(235, 247, 242)
CREAM = RGBColor(248, 249, 246)
WHITE = RGBColor(255, 255, 255)
INK = RGBColor(28, 42, 38)
MUTED = RGBColor(102, 116, 111)
LINE = RGBColor(219, 226, 222)
AMBER = RGBColor(220, 151, 44)
PALE_AMBER = RGBColor(252, 241, 218)
BLUE = RGBColor(72, 100, 182)
PALE_BLUE = RGBColor(236, 240, 251)
RED = RGBColor(183, 66, 56)


def set_background(slide, color=CREAM):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_rect(slide, x, y, w, h, fill, line=None, radius=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line if line else fill
    return shape


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    size=18,
    color=INK,
    bold=False,
    font="Aptos",
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin=0,
):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = margin
    frame.margin_right = margin
    frame.margin_top = margin
    frame.margin_bottom = margin
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_rich_lines(slide, lines, x, y, w, h, size=18, color=INK, bullet_color=GREEN):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(11)
        paragraph.line_spacing = 1.05
        bullet = paragraph.add_run()
        bullet.text = "●  "
        bullet.font.name = "Aptos"
        bullet.font.size = Pt(size - 3)
        bullet.font.color.rgb = bullet_color
        run = paragraph.add_run()
        run.text = line
        run.font.name = "Aptos"
        run.font.size = Pt(size)
        run.font.color.rgb = color
    return box


def add_header(slide, number, kicker, title, subtitle=None, dark=False):
    primary = WHITE if dark else NAVY
    secondary = MINT if dark else GREEN
    add_text(slide, f"{number:02d}", Inches(0.55), Inches(0.43), Inches(0.55), Inches(0.35), 11, secondary, True)
    add_text(slide, kicker.upper(), Inches(1.08), Inches(0.43), Inches(4.8), Inches(0.35), 11, secondary, True)
    add_text(slide, title, Inches(0.55), Inches(0.88), Inches(12.1), Inches(0.7), 30, primary, True)
    if subtitle:
        add_text(slide, subtitle, Inches(0.58), Inches(1.55), Inches(11.8), Inches(0.45), 14, MINT if dark else MUTED)


def add_footer(slide, number, dark=False):
    color = RGBColor(145, 167, 158) if dark else MUTED
    add_text(slide, "STUDENT TASK MANAGER  •  CLOUD COMPUTING MINI PROJECT", Inches(0.58), Inches(7.1), Inches(7.5), Inches(0.2), 8, color, True)
    add_text(slide, str(number), Inches(12.2), Inches(7.08), Inches(0.45), Inches(0.22), 9, color, True, align=PP_ALIGN.RIGHT)


def add_notes(slide, text):
    frame = slide.notes_slide.notes_text_frame
    frame.text = text.strip()


def add_chip(slide, label, x, y, w, fill=PALE_MINT, color=GREEN):
    add_rect(slide, x, y, w, Inches(0.38), fill, fill)
    add_text(slide, label, x, y + Inches(0.015), w, Inches(0.3), 10, color, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def add_metric(slide, value, label, x, y, w, accent=GREEN):
    card = add_rect(slide, x, y, w, Inches(1.15), WHITE, LINE)
    add_rect(slide, x, y, Inches(0.08), Inches(1.15), accent, accent, radius=False)
    add_text(slide, value, x + Inches(0.22), y + Inches(0.17), w - Inches(0.32), Inches(0.45), 24, NAVY, True)
    add_text(slide, label, x + Inches(0.22), y + Inches(0.68), w - Inches(0.32), Inches(0.27), 10, MUTED, True)
    return card


def add_link_button(slide, label, address, x, y, w, fill=GREEN):
    shape = add_rect(slide, x, y, w, Inches(0.55), fill, fill)
    frame = shape.text_frame
    frame.clear()
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = label
    run.font.name = "Aptos"
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = WHITE
    run.hyperlink.address = address
    return shape


def build_deck():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    # 1 — Title
    slide = prs.slides.add_slide(blank)
    set_background(slide, DEEP_NAVY)
    add_rect(slide, Inches(0), Inches(0), Inches(0.22), SLIDE_H, GREEN, GREEN, radius=False)
    add_chip(slide, "LIVE AWS DEMO", Inches(0.7), Inches(0.65), Inches(1.65), MINT, GREEN)
    add_text(slide, "Auto-Scaling &\nSelf-Healing", Inches(0.7), Inches(1.35), Inches(7.5), Inches(1.75), 39, WHITE, True)
    add_text(slide, "Cloud-Native Student Task Manager", Inches(0.73), Inches(3.2), Inches(7.5), Inches(0.55), 22, MINT, True)
    add_text(slide, "Kubernetes  •  Docker  •  Terraform  •  AWS", Inches(0.73), Inches(3.9), Inches(7.2), Inches(0.4), 15, RGBColor(192, 210, 203), True)
    add_rect(slide, Inches(8.65), Inches(0.9), Inches(3.8), Inches(5.35), RGBColor(24, 55, 48), RGBColor(49, 90, 78))
    add_text(slide, "TEAM", Inches(9.05), Inches(1.35), Inches(2.9), Inches(0.3), 11, MINT, True)
    team = [
        ("Myisha Siddiqui", "240905560"),
        ("Manan Bhasin", "240905492"),
        ("Rishabh Raj", "240905420"),
    ]
    top = 1.95
    for idx, (name, reg) in enumerate(team, start=1):
        add_rect(slide, Inches(9.0), Inches(top), Inches(0.48), Inches(0.48), GREEN, GREEN)
        add_text(slide, str(idx), Inches(9.0), Inches(top + 0.04), Inches(0.48), Inches(0.3), 11, WHITE, True, align=PP_ALIGN.CENTER)
        add_text(slide, name, Inches(9.68), Inches(top - 0.01), Inches(2.3), Inches(0.3), 15, WHITE, True)
        add_text(slide, reg, Inches(9.68), Inches(top + 0.3), Inches(2.0), Inches(0.25), 10, RGBColor(176, 198, 190))
        top += 1.13
    add_text(slide, "AWS Sydney  /  ap-southeast-2", Inches(9.05), Inches(5.55), Inches(2.9), Inches(0.3), 10, MINT, True)
    add_footer(slide, 1, dark=True)
    add_notes(slide, """
Good morning. Our project is an auto-scaling and self-healing cloud-native Student Task Manager. We built the application with Spring Boot, packaged it with Docker, provisioned AWS using Terraform, and deployed it on Kubernetes using K3s. The project was completed by Myisha Siddiqui, Manan Bhasin, and Rishabh Raj.
""")

    # 2 — Application and objective
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_header(slide, 2, "The application", "A useful app—and a workload for cloud behaviour")
    add_rect(slide, Inches(0.55), Inches(2.05), Inches(5.05), Inches(4.55), WHITE, LINE)
    add_text(slide, "What the user can do", Inches(0.9), Inches(2.42), Inches(4.2), Inches(0.45), 20, NAVY, True)
    add_rich_lines(
        slide,
        [
            "Create, edit, complete, filter and delete study tasks",
            "Refresh the browser without losing saved task data",
            "Use the same Spring Boot service for the UI and REST API",
            "Persist every task in a private MySQL database",
        ],
        Inches(0.9), Inches(3.05), Inches(4.15), Inches(2.7), 15,
    )
    add_chip(slide, "STATELESS API", Inches(0.9), Inches(5.93), Inches(1.55))
    add_chip(slide, "PERSISTENT DATA", Inches(2.6), Inches(5.93), Inches(1.8), PALE_BLUE, BLUE)
    add_rect(slide, Inches(5.9), Inches(2.05), Inches(6.85), Inches(4.55), WHITE, LINE)
    slide.shapes.add_picture(str(UI_IMAGE), Inches(6.12), Inches(2.27), width=Inches(6.4), height=Inches(3.98))
    add_text(slide, "Live StudySpace browser dashboard", Inches(6.15), Inches(6.29), Inches(4.3), Inches(0.22), 9, MUTED, True)
    add_footer(slide, 2)
    add_notes(slide, """
The visible application is a Student Task Manager called StudySpace. A student can create, edit, complete, filter, and delete tasks. The browser sends requests to the Spring Boot REST API, and the API stores the records in MySQL. Refreshing the page proves that the data is persistent rather than held only in the browser. This real workload gives us something meaningful to scale and recover.
""")

    # 3 — Architecture
    slide = prs.slides.add_slide(blank)
    set_background(slide, DEEP_NAVY)
    add_header(slide, 3, "Architecture", "One request path, one monitoring path", "The database is private; only the browser-facing demo ports are exposed.", dark=True)

    cards = [
        ("BROWSER", "StudySpace UI", PALE_MINT, GREEN),
        ("K3S SERVICE", "NodePort 30080", PALE_BLUE, BLUE),
        ("SPRING BOOT", "1–2 API pods", PALE_AMBER, AMBER),
        ("RDS MYSQL", "Private subnets", RGBColor(247, 232, 231), RED),
    ]
    x_positions = [0.55, 3.75, 6.95, 10.15]
    for index, ((head, sub, fill, accent), x) in enumerate(zip(cards, x_positions)):
        add_rect(slide, Inches(x), Inches(2.35), Inches(2.58), Inches(1.48), WHITE, RGBColor(57, 80, 72))
        add_rect(slide, Inches(x), Inches(2.35), Inches(2.58), Inches(0.12), accent, accent, radius=False)
        add_text(slide, head, Inches(x + 0.22), Inches(2.68), Inches(2.12), Inches(0.32), 13, accent, True)
        add_text(slide, sub, Inches(x + 0.22), Inches(3.12), Inches(2.12), Inches(0.32), 15, NAVY, True)
        if index < len(cards) - 1:
            arrow = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x + 2.67), Inches(2.82), Inches(0.42), Inches(0.48))
            arrow.fill.solid(); arrow.fill.fore_color.rgb = MINT
            arrow.line.color.rgb = MINT

    add_text(slide, "REQUEST + DATA PATH", Inches(0.58), Inches(4.15), Inches(2.3), Inches(0.3), 10, MINT, True)
    add_rect(slide, Inches(0.55), Inches(4.55), Inches(7.9), Inches(1.55), RGBColor(24, 55, 48), RGBColor(49, 90, 78))
    add_text(slide, "Infrastructure delivery", Inches(0.85), Inches(4.85), Inches(2.2), Inches(0.3), 13, MINT, True)
    add_text(slide, "Terraform", Inches(0.85), Inches(5.28), Inches(1.5), Inches(0.3), 16, WHITE, True)
    add_text(slide, "VPC  •  EC2  •  RDS  •  ECR  •  IAM", Inches(2.5), Inches(5.29), Inches(4.8), Inches(0.3), 14, RGBColor(195, 213, 206), True)

    add_rect(slide, Inches(8.75), Inches(4.55), Inches(4.0), Inches(1.55), RGBColor(24, 55, 48), RGBColor(49, 90, 78))
    add_text(slide, "Observability", Inches(9.05), Inches(4.85), Inches(1.7), Inches(0.3), 13, MINT, True)
    add_text(slide, "Prometheus  →  Grafana", Inches(9.05), Inches(5.29), Inches(3.2), Inches(0.3), 16, WHITE, True)
    add_footer(slide, 3, dark=True)
    add_notes(slide, """
The request path is Browser to the Kubernetes NodePort Service, then to one or more Spring Boot pods, and finally to private RDS MySQL. The pods are stateless, so either replica can serve a request while all replicas share the same database. Docker creates the image and ECR stores it. Terraform provisions the VPC, EC2 host, RDS, ECR, IAM, and security rules. Prometheus collects metrics and Grafana visualizes them.
""")

    # 4 — Technology roles
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_header(slide, 4, "Implementation", "How each technology contributes")
    tech = [
        ("Docker", "Builds one portable Spring Boot image", GREEN, PALE_MINT),
        ("Amazon ECR", "Stores the image pulled by K3s", BLUE, PALE_BLUE),
        ("Terraform", "Creates repeatable AWS infrastructure", AMBER, PALE_AMBER),
        ("Kubernetes / K3s", "Runs pods, probes, Services and HPA", GREEN, PALE_MINT),
        ("RDS MySQL", "Keeps task data outside the pods", RED, RGBColor(247, 232, 231)),
        ("GitHub Actions", "Tests the app and validates Terraform", BLUE, PALE_BLUE),
    ]
    positions = [(0.55, 2.05), (4.4, 2.05), (8.25, 2.05), (0.55, 4.25), (4.4, 4.25), (8.25, 4.25)]
    for idx, ((title, body, accent, pale), (x, y)) in enumerate(zip(tech, positions), start=1):
        add_rect(slide, Inches(x), Inches(y), Inches(3.55), Inches(1.75), WHITE, LINE)
        add_rect(slide, Inches(x + 0.25), Inches(y + 0.26), Inches(0.46), Inches(0.46), pale, pale)
        add_text(slide, str(idx), Inches(x + 0.25), Inches(y + 0.31), Inches(0.46), Inches(0.25), 11, accent, True, align=PP_ALIGN.CENTER)
        add_text(slide, title, Inches(x + 0.88), Inches(y + 0.24), Inches(2.3), Inches(0.35), 17, NAVY, True)
        add_text(slide, body, Inches(x + 0.28), Inches(y + 0.88), Inches(2.98), Inches(0.58), 13, MUTED)
    add_footer(slide, 4)
    add_notes(slide, """
Each technology has a clear responsibility. Docker provides a consistent image. ECR stores that image. Terraform makes the infrastructure repeatable. K3s provides Kubernetes features on a cost-conscious EC2 host. RDS keeps data independent of pod lifecycles. GitHub Actions tests the application and validates the infrastructure code. The main design point is separation: pods can be replaced without losing tasks.
""")

    # 5 — Autoscaling
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_header(slide, 5, "Auto-scaling", "HPA responds to CPU demand—and scales back down", "CPU target: 60%  •  Minimum: 1 pod  •  Maximum: 2 pods")
    stages = [
        ("01", "NORMAL", "1 pod", "Low CPU", GREEN),
        ("02", "LOAD", "CPU > 60%", "Peak 319%", AMBER),
        ("03", "SCALE UP", "2 pods", "Both Ready", BLUE),
        ("04", "COOLDOWN", "1 pod", "After 180 s", GREEN),
    ]
    xs = [0.55, 3.75, 6.95, 10.15]
    for index, ((num, label, value, detail, accent), x) in enumerate(zip(stages, xs)):
        add_rect(slide, Inches(x), Inches(2.18), Inches(2.58), Inches(1.82), WHITE, LINE)
        add_text(slide, num, Inches(x + 0.22), Inches(2.42), Inches(0.5), Inches(0.25), 10, accent, True)
        add_text(slide, label, Inches(x + 0.77), Inches(2.4), Inches(1.5), Inches(0.28), 11, accent, True)
        add_text(slide, value, Inches(x + 0.22), Inches(2.93), Inches(2.0), Inches(0.42), 21, NAVY, True)
        add_text(slide, detail, Inches(x + 0.22), Inches(3.5), Inches(2.0), Inches(0.25), 11, MUTED, True)
        if index < len(stages) - 1:
            arrow = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x + 2.67), Inches(2.85), Inches(0.42), Inches(0.48))
            arrow.fill.solid(); arrow.fill.fore_color.rgb = MINT
            arrow.line.color.rgb = MINT

    add_metric(slide, "5,699", "HTTP REQUESTS", Inches(0.55), Inches(4.65), Inches(2.75), GREEN)
    add_metric(slide, "0%", "FAILED REQUESTS", Inches(3.55), Inches(4.65), Inches(2.75), GREEN)
    add_metric(slide, "403.98 ms", "P95 RESPONSE TIME", Inches(6.55), Inches(4.65), Inches(2.75), BLUE)
    add_metric(slide, "43.7/s", "REQUEST RATE", Inches(9.55), Inches(4.65), Inches(2.75), AMBER)
    add_text(slide, "Measured with a bounded k6 run at 30 maximum virtual users.", Inches(0.58), Inches(6.18), Inches(6.5), Inches(0.3), 11, MUTED)
    add_footer(slide, 5)
    add_notes(slide, """
The Horizontal Pod Autoscaler targets 60 percent CPU and is limited to one to two replicas for this low-cost demo. During a bounded k6 load run, CPU rose above the target and peaked at 319 percent of the requested CPU. Kubernetes increased the deployment from one pod to two, and both became Ready. The run completed 5,699 requests with zero failures and a p95 of 403.98 milliseconds. After the 180-second stabilization period, it returned to one pod.
""")

    # 6 — Self-healing
    slide = prs.slides.add_slide(blank)
    set_background(slide, DEEP_NAVY)
    add_header(slide, 6, "Self-healing", "Kubernetes restores the desired state", "Two controlled tests proved replacement and health-probe recovery.", dark=True)
    add_rect(slide, Inches(0.55), Inches(2.13), Inches(5.92), Inches(3.95), RGBColor(24, 55, 48), RGBColor(49, 90, 78))
    add_chip(slide, "TEST 1", Inches(0.88), Inches(2.48), Inches(0.95), MINT, GREEN)
    add_text(slide, "Pod deletion", Inches(0.88), Inches(3.02), Inches(4.5), Inches(0.45), 23, WHITE, True)
    add_rich_lines(
        slide,
        [
            "Running API pod was deleted",
            "Deployment detected a missing replica",
            "Replacement pod reached 1/1 Running",
            "Deployment remained available",
        ],
        Inches(0.88), Inches(3.68), Inches(4.9), Inches(1.9), 14, RGBColor(222, 236, 230), MINT,
    )
    add_rect(slide, Inches(6.87), Inches(2.13), Inches(5.88), Inches(3.95), RGBColor(24, 55, 48), RGBColor(49, 90, 78))
    add_chip(slide, "TEST 2", Inches(7.2), Inches(2.48), Inches(0.95), PALE_AMBER, AMBER)
    add_text(slide, "Liveness failure", Inches(7.2), Inches(3.02), Inches(4.5), Inches(0.45), 23, WHITE, True)
    add_rich_lines(
        slide,
        [
            "Java process was paused intentionally",
            "Liveness probe failed three times",
            "Kubernetes restarted the container",
            "Readiness returned UP",
        ],
        Inches(7.2), Inches(3.68), Inches(4.85), Inches(1.9), 14, RGBColor(222, 236, 230), MINT,
    )
    add_rect(slide, Inches(0.55), Inches(6.34), Inches(12.2), Inches(0.42), RGBColor(52, 72, 65), RGBColor(52, 72, 65))
    add_text(slide, "Scope: pod/container recovery on one EC2 node — not EC2 host failover.", Inches(0.75), Inches(6.41), Inches(11.8), Inches(0.22), 10, RGBColor(207, 223, 216), True, align=PP_ALIGN.CENTER)
    add_footer(slide, 6, dark=True)
    add_notes(slide, """
Self-healing was demonstrated in two ways. First, we deleted the running API pod. The Deployment controller compared the actual state with the desired state of one replica and created a replacement, which reached 1/1 Running. Second, we intentionally paused the Java process. After the liveness probe failed, Kubernetes restarted the container and readiness returned UP. This proves pod and container recovery; it does not claim EC2 host-level failover because the cluster has one node.
""")

    # 7 — Monitoring
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_header(slide, 7, "Monitoring", "Prometheus collects; Grafana explains")
    add_rect(slide, Inches(0.55), Inches(2.03), Inches(4.0), Inches(4.65), NAVY, NAVY)
    add_text(slide, "LIVE PIPELINE", Inches(0.9), Inches(2.42), Inches(2.0), Inches(0.25), 10, MINT, True)
    pipeline = [
        ("1", "Spring Boot Actuator", "API + JVM metrics"),
        ("2", "Prometheus", "Scrape every 15 seconds"),
        ("3", "Grafana", "Six-panel dashboard"),
    ]
    y = 3.0
    for num, title, sub in pipeline:
        add_rect(slide, Inches(0.9), Inches(y), Inches(0.5), Inches(0.5), GREEN, GREEN)
        add_text(slide, num, Inches(0.9), Inches(y + 0.05), Inches(0.5), Inches(0.27), 11, WHITE, True, align=PP_ALIGN.CENTER)
        add_text(slide, title, Inches(1.63), Inches(y - 0.02), Inches(2.1), Inches(0.3), 15, WHITE, True)
        add_text(slide, sub, Inches(1.63), Inches(y + 0.32), Inches(2.0), Inches(0.3), 10, RGBColor(182, 204, 196))
        y += 1.1
    add_text(slide, "Prometheus targets: API, Kubernetes state and Prometheus itself", Inches(0.9), Inches(6.05), Inches(3.05), Inches(0.42), 10, MINT, True)

    add_text(slide, "What the dashboard shows", Inches(4.95), Inches(2.12), Inches(4.2), Inches(0.4), 20, NAVY, True)
    panels = [
        ("Request rate", "traffic"),
        ("Error rate", "reliability"),
        ("API latency", "performance"),
        ("Ready replicas", "availability"),
        ("JVM heap", "runtime"),
        ("Process CPU", "scaling input"),
    ]
    positions = [(4.95, 2.78), (8.8, 2.78), (4.95, 3.85), (8.8, 3.85), (4.95, 4.92), (8.8, 4.92)]
    for (title, sub), (x, y) in zip(panels, positions):
        add_rect(slide, Inches(x), Inches(y), Inches(3.35), Inches(0.85), WHITE, LINE)
        add_rect(slide, Inches(x + 0.18), Inches(y + 0.2), Inches(0.1), Inches(0.45), GREEN, GREEN, radius=False)
        add_text(slide, title, Inches(x + 0.45), Inches(y + 0.13), Inches(2.45), Inches(0.25), 13, NAVY, True)
        add_text(slide, sub.upper(), Inches(x + 0.45), Inches(y + 0.47), Inches(2.3), Inches(0.2), 8, MUTED, True)
    add_rect(slide, Inches(4.95), Inches(6.05), Inches(7.2), Inches(0.5), PALE_MINT, PALE_MINT)
    add_text(slide, "Grafana health: database OK  •  Version 12.4.13  •  Dashboard provisioned from code", Inches(5.15), Inches(6.17), Inches(6.8), Inches(0.22), 10, GREEN, True, align=PP_ALIGN.CENTER)
    add_footer(slide, 7)
    add_notes(slide, """
Spring Boot Actuator exposes application and JVM metrics. Prometheus scrapes those metrics every 15 seconds together with Kubernetes state metrics. Grafana presents six operational panels: request rate, error rate, latency, ready replicas, JVM heap, and process CPU. The application also has a Demo Console that combines current Kubernetes and Prometheus information on one page. Grafana health is currently OK and the dashboard is provisioned from code.
""")

    # 8 — Demo flow
    slide = prs.slides.add_slide(blank)
    set_background(slide, DEEP_NAVY)
    add_header(slide, 8, "Live demonstration", "Show the result in four clean steps", "All three links are clickable in Slide Show mode.", dark=True)
    steps = [
        ("1", "Application", "Create a task → complete it → refresh", "http://32.236.242.165:30080/"),
        ("2", "Demo Console", "Show readiness, pod, HPA and metrics", "http://32.236.242.165:30080/demo.html"),
        ("3", "Autoscaling", "Run bounded load and watch replicas 1 → 2", "http://32.236.242.165:30080/demo.html"),
        ("4", "Grafana", "Open the live six-panel operations view", "http://32.236.242.165:30300/d/student-task-manager-ops/student-task-manager-operations"),
    ]
    y = 2.15
    for num, title, detail, url in steps:
        add_rect(slide, Inches(0.58), Inches(y), Inches(12.15), Inches(0.9), RGBColor(24, 55, 48), RGBColor(49, 90, 78))
        add_rect(slide, Inches(0.82), Inches(y + 0.2), Inches(0.48), Inches(0.48), GREEN, GREEN)
        add_text(slide, num, Inches(0.82), Inches(y + 0.25), Inches(0.48), Inches(0.25), 11, WHITE, True, align=PP_ALIGN.CENTER)
        add_text(slide, title, Inches(1.55), Inches(y + 0.17), Inches(1.65), Inches(0.3), 15, WHITE, True)
        add_text(slide, detail, Inches(3.35), Inches(y + 0.18), Inches(5.8), Inches(0.3), 13, RGBColor(207, 223, 216))
        add_link_button(slide, "OPEN  ↗", url, Inches(10.55), Inches(y + 0.18), Inches(1.55), GREEN)
        y += 1.05

    add_rect(slide, Inches(0.58), Inches(6.48), Inches(12.15), Inches(0.39), RGBColor(75, 51, 39), RGBColor(75, 51, 39))
    add_text(slide, "Demo access is temporary and unauthenticated. RDS remains private; close public ports after presenting.", Inches(0.82), Inches(6.56), Inches(11.65), Inches(0.2), 10, RGBColor(252, 221, 185), True, align=PP_ALIGN.CENTER)
    add_footer(slide, 8, dark=True)
    add_notes(slide, """
For the live demonstration, first open StudySpace. Create a task, mark it complete, and refresh to prove that RDS persists the data. Next open the Demo Console and point out readiness, the Deployment, pod state, HPA configuration, and live metrics. Start the bounded read-only load generator and explain that the HPA can move from one to two replicas when CPU exceeds 60 percent. Finally open Grafana and show the six panels. Conclude that Docker delivered the image, Terraform built AWS, Kubernetes scaled and recovered the workload, and Prometheus plus Grafana made the behaviour visible. Public access is temporary and should be closed after the presentation.
""")

    # Set core document metadata.
    prs.core_properties.title = "Auto-Scaling and Self-Healing Cloud-Native Student Task Manager"
    prs.core_properties.subject = "CCD mini project presentation"
    prs.core_properties.author = "Myisha Siddiqui, Manan Bhasin, Rishabh Raj"
    prs.core_properties.comments = "Generated from verified project evidence."
    prs.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_deck())
